"""Excel レポート（概要・一覧・カテゴリ別・評価別・前回との差分）を作る。"""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from webcollect.diff import Change
from webcollect.parser import Book

HEAD = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
KIND_FILL = {"新規": "E2EFDA", "掲載終了": "EDEDED", "値上げ": "FCE4D6", "値下げ": "DDEBF7", "在庫切れ": "F8CBAD", "再入荷": "C6EFCE"}


def _header(ws, row: int, values: list[str]) -> None:
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill, c.font = HEAD, HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")


def _style_chart(ch) -> None:
    # 新しい Excel では軸が隠れる初期値になっているので、明示して表示する。棒は1色にそろえる
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.varyColors = False
    ch.legend = None
    ch.series[0].graphicalProperties.solidFill = "2F6EBA"
    ch.series[0].graphicalProperties.line.solidFill = "2F6EBA"
    if ch.title is not None:
        ch.title.overlay = False  # タイトルをグラフの描画域に重ねない


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_report(books: list[Book], changes: list[Change] | None, path: Path, *, source: str, run_at: datetime,
                 has_previous: bool) -> Path:
    wb = Workbook()

    # 概要
    ws = wb.active
    ws.title = "概要"
    ws["A1"] = "Web データ収集レポート"
    ws["A1"].font = Font(bold=True, size=16)
    rows = [("取得日時", run_at.strftime("%Y-%m-%d %H:%M")), ("取得元", source), ("件数", len(books)),
            ("カテゴリ数", len({b.category for b in books})),
            ("在庫ありの件数", sum(b.in_stock for b in books)),
            ("平均価格（£）", float(sum((b.price for b in books), Decimal(0)) / len(books)) if books else 0)]
    if has_previous and changes is not None:
        for k in ("新規", "掲載終了", "値上げ", "値下げ", "在庫切れ", "再入荷"):
            rows.append((f"前回からの{k}", sum(c.kind == k for c in changes)))
    else:
        rows.append(("前回との比較", "初回のため比較なし"))
    for i, (k, v) in enumerate(rows, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
    ws["B8"].number_format = "#,##0.00"
    _widths(ws, [22, 50])

    # 一覧
    ws = wb.create_sheet("一覧")
    detail = any(b.stock_count is not None or b.upc for b in books)
    cols = ["書名", "カテゴリ", "価格（£）", "評価（5段階）", "在庫"] + (["在庫数", "商品コード"] if detail else []) + ["URL"]
    _header(ws, 1, cols)
    for r, b in enumerate(books, 2):
        row = [b.title, b.category, float(b.price), b.rating, "あり" if b.in_stock else "なし"]
        row += [b.stock_count, b.upc] if detail else []
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
        ws.cell(row=r, column=3).number_format = "#,##0.00"
        link = ws.cell(row=r, column=len(cols), value=b.url)
        link.hyperlink, link.style = b.url, "Hyperlink"
    if books:
        t = Table(displayName="BookList", ref=f"A1:{get_column_letter(len(cols))}{len(books) + 1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(t)
    ws.freeze_panes = "A2"
    _widths(ws, [48, 20, 13, 16, 9] + ([11, 20] if detail else []) + [60])

    # カテゴリ別
    ws = wb.create_sheet("カテゴリ別")
    by = defaultdict(list)
    for b in books:
        by[b.category].append(b)
    _header(ws, 1, ["カテゴリ", "件数", "平均価格（£）", "最安（£）", "最高（£）", "在庫ありの件数"])
    stats = sorted(by.items(), key=lambda kv: (-len(kv[1]), kv[0]))
    for r, (cat, bs) in enumerate(stats, 2):
        prices = [float(b.price) for b in bs]
        ws.append([cat, len(bs), round(sum(prices) / len(prices), 2), min(prices), max(prices), sum(b.in_stock for b in bs)])
        for c in (3, 4, 5):
            ws.cell(row=r, column=c).number_format = "#,##0.00"
    _widths(ws, [22, 8, 14, 11, 11, 14])
    ws.freeze_panes = "A2"
    if stats:
        n = min(15, len(stats))
        ch = BarChart()
        ch.type, ch.title = "bar", f"件数の多いカテゴリ（上位{n}）"
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=n + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
        ch.y_axis.scaling.orientation = "minMax"
        ch.x_axis.scaling.orientation = "maxMin"  # 多い順に上から並べる
        ch.y_axis.crosses = "max"  # 上下を入れ替えても、目盛りは下に出す
        _style_chart(ch)
        ch.height, ch.width = 11, 18
        ws.add_chart(ch, "H2")

    # 評価別
    ws = wb.create_sheet("評価別")
    _header(ws, 1, ["評価", "件数", "平均価格（£）"])
    for star in range(1, 6):
        bs = [b for b in books if b.rating == star]
        ws.append([f"★{star}", len(bs), round(sum(float(b.price) for b in bs) / len(bs), 2) if bs else 0])
        ws.cell(row=star + 1, column=3).number_format = "#,##0.00"
    _widths(ws, [10, 8, 14])
    ch = BarChart()
    ch.title = "評価ごとの件数"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
    _style_chart(ch)
    ch.height, ch.width = 8, 14
    ws.add_chart(ch, "E2")

    # 前回との差分
    ws = wb.create_sheet("前回との差分")
    _header(ws, 1, ["種類", "書名", "前回", "今回", "URL"])
    if not has_previous or changes is None:
        ws["A2"] = "初回のため、比較する前回のデータがありません"
    elif not changes:
        ws["A2"] = "前回から変わったところはありません"
    for r, c in enumerate(changes or [], 2):
        ws.append([c.kind, c.title, c.before, c.after, c.url])
        ws.cell(row=r, column=5).hyperlink = c.url
        ws.cell(row=r, column=5).style = "Hyperlink"
        if c.kind in KIND_FILL:  # 種類ごとに色を付ける（並べ替えても色が付いたままになるよう、セルに直接）
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=KIND_FILL[c.kind])
    ws.freeze_panes = "A2"
    _widths(ws, [12, 48, 12, 12, 60])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
