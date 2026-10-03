import csv

from openpyxl import load_workbook

import webcollect.cli as cli
from conftest import FakeFetcher
from webcollect.config import ConfigError, load_config

import pytest


def write_cfg(tmp_path, extra=""):
    p = tmp_path / "config.yaml"
    p.write_text(f"categories: [Mystery]\ndelay_seconds: 1.0\n{extra}", encoding="utf-8")
    return p


@pytest.mark.parametrize("text,msg", [
    ("delay_seconds: 0.1\n", "0.5 以上"),
    ("base_url: ftp://x/\n", "http"),
    ("typo_key: 1\n", "知らない設定項目"),
    ("- a\n- b\n", "キー: 値"),
])
def test_config_errors(tmp_path, text, msg):
    p = tmp_path / "c.yaml"
    p.write_text(text, encoding="utf-8")
    with pytest.raises(ConfigError, match=msg):
        load_config(p)


def test_config_missing(tmp_path):
    with pytest.raises(ConfigError, match="見つかりません"):
        load_config(tmp_path / "none.yaml")


def test_cli_bad_config_returns_2(tmp_path, capsys):
    assert cli.main(["--config", str(tmp_path / "none.yaml")]) == 2
    assert "設定エラー" in capsys.readouterr().err


def test_cli_two_runs_make_diff(tmp_path, monkeypatch):
    monkeypatch.setattr(cli, "Fetcher", lambda *a, **k: FakeFetcher())
    cfg = write_cfg(tmp_path)
    assert cli.main(["--config", str(cfg), "--no-notify"]) == 0
    out = tmp_path / "output"
    first = sorted(out.glob("report_*.xlsx"))
    assert len(first) == 1 and "初回" in load_workbook(first[0])["前回との差分"]["A2"].value

    # 前回分の価格を1件だけ書き換えて、2回目で「値上げ」として出るか確かめる
    snap = sorted((out / "history").glob("snapshot_*.csv"))[0]
    rows = list(csv.DictReader(snap.open(encoding="utf-8-sig")))
    rows[0]["price"] = "1.00"
    with snap.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=rows[0].keys())
        w.writeheader()
        w.writerows(rows)
    import time
    time.sleep(1.1)  # 履歴のファイル名（秒単位）が重ならないように
    monkeypatch.setattr(cli, "datetime", type("D", (), {"now": staticmethod(lambda: __import__("datetime").datetime(2099, 1, 1, 0, 1))}))
    assert cli.main(["--config", str(cfg), "--no-notify"]) == 0
    second = out / "report_20990101_0001.xlsx"
    ws = load_workbook(second)["前回との差分"]
    assert ws["A2"].value == "値上げ" and ws["C2"].value == "£1.00"


def test_cli_stops_on_unknown_category(tmp_path, monkeypatch):
    monkeypatch.setattr(cli, "Fetcher", lambda *a, **k: FakeFetcher())
    cfg = write_cfg(tmp_path)
    assert cli.main(["--config", str(cfg), "--categories", "ないカテゴリ", "--no-notify"]) == 1
